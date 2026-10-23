"""Monta o artigo completo em .docx a partir de artigo_agrotec.md.

Modelo: Revista Agropecuária Técnica (AGROTEC/UFPB), A4, margens de 2,0 cm
no topo e na base e 2,5 cm nas laterais, Times New Roman 12, entrelinhas 1,5,
linhas e páginas numeradas, seções numeradas, tabelas sem linhas verticais,
referências em ordem alfabética.

Uso:
    python artigo_cientifico/gera_docx_agrotec.py
"""

from pathlib import Path
import subprocess

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

PASTA = Path(__file__).resolve().parent
ORIGEM = PASTA / "artigo_agrotec.md"
SAIDA = PASTA / "RECUPERA-SOLO - Artigo AGROTEC.docx"

SECOES = {"Resumo", "Abstract", "1 Introdução", "2 Material e Métodos",
          "3 Resultados e Discussão", "4 Conclusão", "Referências"}
FONTE = "Times New Roman"


def formatar_fonte(run, tamanho, negrito=None):
    run.font.name = FONTE
    run.font.size = Pt(tamanho)
    if negrito is not None:
        run.font.bold = negrito
    rpr = run._element.get_or_add_rPr()
    fontes = rpr.find(qn("w:rFonts"))
    if fontes is None:
        fontes = OxmlElement("w:rFonts")
        rpr.append(fontes)
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fontes.set(qn(atributo), FONTE)


def formatar_paragrafo(paragrafo, tamanho=12, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY,
                       recuo=1.0, antes=0, depois=0, negrito=None, juntar=False):
    formato = paragrafo.paragraph_format
    formato.alignment = alinhamento
    formato.first_line_indent = Cm(recuo)
    formato.space_before = Pt(antes)
    formato.space_after = Pt(depois)
    formato.keep_with_next = juntar
    formato.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    for run in paragrafo.runs:
        formatar_fonte(run, tamanho, negrito)


def captions(paragrafo, palavra, tamanho=10):
    """Reescreve a legenda com o prefixo em negrito."""
    texto = paragrafo.text.strip()
    resto = texto[len(palavra):].lstrip(" -")
    for run in list(paragrafo.runs):
        run._element.getparent().remove(run._element)
    inicio = paragrafo.add_run(palavra)
    formatar_fonte(inicio, tamanho, True)
    if resto:
        corpo = paragrafo.add_run(" - " + resto)
        formatar_fonte(corpo, tamanho, False)


def bordas_tabela(tabela):
    """Aplica apenas linhas horizontais, sem linhas verticais."""
    propriedades = tabela._tbl.tblPr
    bordas = OxmlElement("w:tblBorders")
    for lado, espessura in (("top", 8), ("bottom", 8), ("insideH", 4)):
        linha = OxmlElement(f"w:{lado}")
        linha.set(qn("w:val"), "single")
        linha.set(qn("w:sz"), str(espessura))
        linha.set(qn("w:space"), "0")
        linha.set(qn("w:color"), "000000")
        bordas.append(linha)
    for lado in ("left", "right", "insideV"):
        linha = OxmlElement(f"w:{lado}")
        linha.set(qn("w:val"), "nil")
        bordas.append(linha)
    propriedades.append(bordas)


def ajustar_tabela(tabela, larguras, numericas=()):
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.autofit = False
    bordas_tabela(tabela)

    total = int(sum(largura.cm for largura in larguras) / 2.54 * 1440)
    propriedades = tabela._tbl.tblPr
    largura_tabela = OxmlElement("w:tblW")
    largura_tabela.set(qn("w:w"), str(total))
    largura_tabela.set(qn("w:type"), "dxa")
    propriedades.append(largura_tabela)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    propriedades.append(layout)
    for grade, largura in zip(tabela._tbl.tblGrid.gridCol_lst, larguras):
        grade.w = largura
    for indice, coluna in enumerate(tabela.columns):
        coluna.width = larguras[indice]
    for numero, linha in enumerate(tabela.rows):
        for indice, celula in enumerate(linha.cells):
            celula.width = larguras[indice]
            for paragrafo in celula.paragraphs:
                alinhamento = WD_ALIGN_PARAGRAPH.RIGHT if indice in numericas else WD_ALIGN_PARAGRAPH.LEFT
                if numero == 0:
                    alinhamento = WD_ALIGN_PARAGRAPH.CENTER
                formatar_paragrafo(paragrafo, tamanho=10, alinhamento=alinhamento,
                                   recuo=0, antes=1, depois=1)
            if numero == 0:
                for paragrafo in celula.paragraphs:
                    for run in paragrafo.runs:
                        formatar_fonte(run, 10, True)

    for numero, linha in enumerate(tabela.rows):
        propriedades = linha._tr.get_or_add_trPr()
        propriedades.append(OxmlElement("w:cantSplit"))
        if numero == 0:
            propriedades.append(OxmlElement("w:tblHeader"))
        if numero < len(tabela.rows) - 1:
            for celula in linha.cells:
                for paragrafo in celula.paragraphs:
                    paragrafo.paragraph_format.keep_with_next = True


def numerar_paginas(secao):
    rodape = secao.footer.paragraphs[0]
    rodape.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    campo = OxmlElement("w:fldSimple")
    campo.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    texto = OxmlElement("w:t")
    texto.text = "1"
    run.append(texto)
    campo.append(run)
    rodape._p.append(campo)
    for run in rodape.runs:
        formatar_fonte(run, 10)


def numerar_linhas(secao):
    linhas = OxmlElement("w:lnNumType")
    linhas.set(qn("w:countBy"), "1")
    linhas.set(qn("w:restart"), "continuous")
    linhas.set(qn("w:distance"), "360")
    secao._sectPr.append(linhas)


def ajustar_figuras(documento):
    larguras = [Cm(16), Cm(8)]
    for indice, forma in enumerate(documento.inline_shapes):
        largura = larguras[indice] if indice < len(larguras) else Cm(17)
        proporcao = forma.height / forma.width
        forma.width = largura
        forma.height = int(largura * proporcao)
    for paragrafo in documento.paragraphs:
        if paragrafo._p.findall(".//" + qn("w:drawing")):
            formatar_paragrafo(paragrafo, alinhamento=WD_ALIGN_PARAGRAPH.CENTER, recuo=0)


def montar():
    subprocess.run(["pandoc", ORIGEM.name, "-o", str(SAIDA)], check=True, cwd=PASTA)
    documento = Document(SAIDA)

    secao = documento.sections[0]
    secao.page_width = Cm(21)
    secao.page_height = Cm(29.7)
    secao.top_margin = Cm(2.0)
    secao.bottom_margin = Cm(2.0)
    secao.left_margin = Cm(2.5)
    secao.right_margin = Cm(2.5)
    secao.header_distance = Cm(1.0)
    secao.footer_distance = Cm(1.0)

    normal = documento.styles["Normal"]
    normal.font.name = FONTE
    normal.font.size = Pt(12)

    em_referencias = False
    apos_resumo = False

    for indice, paragrafo in enumerate(documento.paragraphs):
        texto = paragrafo.text.strip()
        if indice == 0:
            formatar_paragrafo(paragrafo, alinhamento=WD_ALIGN_PARAGRAPH.CENTER,
                               recuo=0, negrito=True, depois=6)
            continue
        if indice == 1:
            formatar_paragrafo(paragrafo, alinhamento=WD_ALIGN_PARAGRAPH.CENTER,
                               recuo=0, depois=2)
            continue
        if texto.startswith("1 Universidade Federal"):
            formatar_paragrafo(paragrafo, tamanho=10, alinhamento=WD_ALIGN_PARAGRAPH.CENTER,
                               recuo=0, depois=12)
            continue
        if texto == "RECUPERA-SOLO: A DECISION SUPPORT SYSTEM FOR SODIC SOIL RECLAMATION":
            formatar_paragrafo(paragrafo, alinhamento=WD_ALIGN_PARAGRAPH.CENTER,
                               recuo=0, negrito=True, antes=12, depois=6)
            continue
        if texto in SECOES:
            formatar_paragrafo(paragrafo, alinhamento=WD_ALIGN_PARAGRAPH.CENTER,
                               recuo=0, negrito=True, antes=12, depois=6, juntar=True)
            em_referencias = texto == "Referências"
            apos_resumo = texto in {"Resumo", "Abstract"}
            continue
        if texto[:3] in ("2.1", "2.2", "2.3", "2.4", "2.5", "2.6",
                         "3.1", "3.2", "3.3", "3.4"):
            formatar_paragrafo(paragrafo, alinhamento=WD_ALIGN_PARAGRAPH.LEFT,
                               recuo=0, negrito=True, antes=6, depois=3, juntar=True)
            continue
        if texto.startswith("Tabela"):
            formatar_paragrafo(paragrafo, tamanho=10, alinhamento=WD_ALIGN_PARAGRAPH.CENTER,
                               recuo=0, antes=6, depois=3, juntar=True)
            captions(paragrafo, texto.split(" - ")[0])
            continue
        if texto.startswith("Figura"):
            formatar_paragrafo(paragrafo, tamanho=10, alinhamento=WD_ALIGN_PARAGRAPH.CENTER,
                               recuo=0, antes=3, depois=3)
            captions(paragrafo, texto.split(" - ")[0])
            continue
        if texto.startswith("Fonte:"):
            formatar_paragrafo(paragrafo, tamanho=10, alinhamento=WD_ALIGN_PARAGRAPH.LEFT,
                               recuo=0, depois=6)
            continue
        if em_referencias:
            formatar_paragrafo(paragrafo, alinhamento=WD_ALIGN_PARAGRAPH.LEFT,
                               recuo=0, depois=6)
            continue
        if texto.startswith("Palavras-chave") or texto.startswith("Keywords"):
            formatar_paragrafo(paragrafo, recuo=0, antes=6, depois=6,
                               alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
            rotulo = "Palavras-chave:" if texto.startswith("Palavras") else "Keywords:"
            resto = texto[len(rotulo):].strip()
            for run in list(paragrafo.runs):
                run._element.getparent().remove(run._element)
            inicio = paragrafo.add_run(rotulo + " ")
            formatar_fonte(inicio, 12, True)
            corpo = paragrafo.add_run(resto)
            formatar_fonte(corpo, 12, False)
            continue
        if apos_resumo:
            formatar_paragrafo(paragrafo, recuo=0)
            continue
        formatar_paragrafo(paragrafo)

    ajustar_tabela(documento.tables[0], [Cm(6.4), Cm(3.8), Cm(2.9), Cm(2.9)], numericas=(2, 3))
    ajustar_tabela(documento.tables[1], [Cm(3.0), Cm(3.4), Cm(1.8), Cm(1.8),
                                         Cm(1.8), Cm(2.0), Cm(2.2)], numericas=(2, 3, 4, 5, 6))
    ajustar_tabela(documento.tables[2], [Cm(5.0), Cm(5.5), Cm(5.5)])
    ajustar_figuras(documento)
    numerar_paginas(secao)
    numerar_linhas(secao)

    propriedades = documento.core_properties
    propriedades.title = "RECUPERA-SOLO: sistema de apoio à decisão para recuperação de solos sódicos"
    propriedades.subject = "Artigo científico"
    propriedades.author = "Nildo Dias; Francisco de Assis P. V. de Arruda"

    documento.save(SAIDA)
    return SAIDA


if __name__ == "__main__":
    print(montar())
